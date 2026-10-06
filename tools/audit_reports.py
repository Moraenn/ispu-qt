import argparse
import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader

parser=argparse.ArgumentParser()
parser.add_argument('--original',type=Path,required=True)
parser.add_argument('--rendered',type=Path,required=True)
args=parser.parse_args()
root=Path(__file__).resolve().parents[1]
report={}
with ZipFile(args.original) as archive:
    for path in sorted((root/'reports/docx').glob('*.docx')):
        member=next(n for n in archive.namelist() if n.replace('\\','/').endswith('/reports/docx/'+path.name))
        original=archive.read(member)
        old=Document(io.BytesIO(original)); new=Document(path)
        with ZipFile(io.BytesIO(original)) as baseline, ZipFile(path) as changed:
            preserved={name:baseline.read(name)==changed.read(name) for name in
                       ['word/styles.xml','word/numbering.xml','word/header1.xml','word/footer1.xml'] if name in baseline.namelist()}
        assert all(preserved.values()), (path,preserved)
        for a,b in zip(old.sections,new.sections):
            for property in ['page_width','page_height','left_margin','right_margin','top_margin','bottom_margin','header_distance','footer_distance','different_first_page_header_footer']:
                assert getattr(a,property)==getattr(b,property),(path,property)
        expected=31 if path.name.startswith('Ismaylov') else {1:3,2:5,3:3,4:7,5:13}[int(path.name[2])]
        assert len(new.inline_shapes)==expected,(path,len(new.inline_shapes),expected)
        text='\n'.join(p.text for p in new.paragraphs if p.style.name!='Listing')
        assert 'Скриншоты работающих окон пока не получены' not in text
        reader=PdfReader(args.rendered/path.stem/(path.stem+'.pdf'))
        pages=[p.extract_text() or '' for p in reader.pages]
        report[path.name]={'pages':len(pages),'inline_images':len(new.inline_shapes),'preserved_parts':preserved,
                           'empty_pages':[i+1 for i,t in enumerate(pages) if not t.strip()],
                           'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
        sheetdir=args.rendered/'sheets'/path.stem
        sheetdir.mkdir(parents=True,exist_ok=True)
        images=sorted((args.rendered/path.stem).glob('page-*.png'),key=lambda p:int(p.stem.split('-')[1]))
        for first in range(0,len(images),4):
            with Image.open(images[first]) as im: width,height=im.size
            sheet=Image.new('RGB',(2*width,2*(height+28)), '#cccccc')
            draw=ImageDraw.Draw(sheet)
            for offset,png in enumerate(images[first:first+4]):
                x=(offset%2)*width; y=(offset//2)*(height+28)
                draw.text((x+10,y+7),f'{path.stem} page {first+offset+1}',fill='black')
                with Image.open(png) as im: sheet.paste(im,(x,y+28))
            sheet.save(sheetdir/f'sheet-{first//4+1:02}.png')
(root/'verification/report-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
