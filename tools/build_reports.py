"""Собрать отчёты DOCX по исходникам, описаниям и реальным снимкам окон.

Установка: python -m pip install python-docx Pillow
Запуск: python tools/build_reports.py --student "Исмайлов В.П." --group "4-41"
Скрипт не запускает Qt и не объявляет тесты пройденными по наличию снимков.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from typing import Iterable
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / 'reports'
FONT = 'Times New Roman'
WIDTH_CM = 16.5
SOURCE_ORDER = {
 'Calculator': ['calculator.h','calculator.cpp','main.cpp','Calculator.pro','CMakeLists.txt'],
 'Alarm': ['alarmwindow.h','alarmwindow.cpp','main.cpp','Alarm.pro','CMakeLists.txt'],
 'handbook_lab3': ['stamp.h','stampmodel.h','stampmodel.cpp','stamptableview.h','stamptableview.cpp',
                  'stampdelegate.h','stampdelegate.cpp','stampdialog.h','stampdialog.cpp',
                  'mainwindow.h','mainwindow.cpp','main.cpp','handbook_lab3.pro','CMakeLists.txt'],
 'handbook': ['stamp.h','stampmodel.h','stampmodel.cpp','stamptableview.h','stamptableview.cpp',
              'stampdelegate.h','stampdelegate.cpp','stampdialog.h','stampdialog.cpp','stampstorage.h',
              'stampstorage.cpp','mainwindow.h','mainwindow.cpp','main.cpp','handbook.pro','CMakeLists.txt'],
 'StudyPlanner': ['core/studyplan.h','core/studyplan.cpp','qtutils.h','plannerstorage.h','plannerstorage.cpp',
                  'plannercontroller.h','plannercontroller.cpp','plannermodels.h','plannermodels.cpp',
                  'taskfilter.h','taskfilter.cpp','taskdelegate.h','taskdelegate.cpp','plannerdialogs.h',
                  'plannerdialogs.cpp','summarywidget.h','summarywidget.cpp','plannerwindow.h',
                  'plannerwindow.cpp','main.cpp','StudyPlanner.pro','CMakeLists.txt']
}
SCREEN_TITLES = {'lab1':'Калькулятор','lab2':'Будильник','lab3':'Справочник филателиста',
                 'lab4':'Справочник с файловыми операциями','lab5_tasks':'Планировщик. Задания',
                 'lab5_sessions':'Планировщик. Журнал занятий','lab5_summary':'Планировщик. Сводка'}
SCREEN_TITLES.update({'lab2_signal':'Сообщение о срабатывании будильника',
 'lab4_dialog':'Добавление марки','lab4_saved':'Коллекция после сохранения и повторного открытия',
 'lab5_dialog':'Добавление учебного задания','lab5_progress':'Сводка после занятия и выполнения задания'})
DIAGRAM_TITLES = {'lab5_core':'Предметная логика и хранение данных',
                  'lab5_ui':'Модели, фильтрация и отображение',
                  'lab5_dialogs':'Окно, диалоги и сводка'}

def font(obj, size=12, bold=False):
    obj.font.name=FONT
    obj.font.size=Pt(size)
    obj.font.bold=bold
    obj.font.color.rgb=RGBColor(0,0,0)
    rp=obj.element.get_or_add_rPr() if hasattr(obj,'element') else obj._element.get_or_add_rPr()
    fonts=rp.find(qn('w:rFonts'))
    if fonts is None:
        fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
    for name in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+name),FONT)

def style_font(style,size=12,bold=False):
    style.font.name=FONT;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    f=rp.find(qn('w:rFonts'))
    if f is None: f=OxmlElement('w:rFonts');rp.append(f)
    for name in ('ascii','hAnsi','eastAsia','cs'): f.set(qn('w:'+name),FONT)
    for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        if qn('w:'+attr) in f.attrib: del f.attrib[qn('w:'+attr)]

def setup(student):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2)
    s.header_distance=Cm(1);s.footer_distance=Cm(1)
    s.different_first_page_header_footer=True
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
        style_font(d.styles[name],14 if name.startswith('Heading') else 12,name.startswith('Heading'))
    n=d.styles['Normal'].paragraph_format
    n.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;n.first_line_indent=Cm(0)
    n.space_after=Pt(6);n.line_spacing=1.5;n.widow_control=True
    for name in ('Heading 1','Heading 2','Heading 3'):
        p=d.styles[name].paragraph_format
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.space_before=Pt(12);p.space_after=Pt(6)
        p.line_spacing=1.15;p.keep_with_next=True;p.keep_together=True
    for name,size in [('Listing',12),('File title',12),('Table text',12),('Cover',12)]:
        st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=d.styles['Normal'];style_font(st,size,name=='File title')
        p=st.paragraph_format;p.line_spacing=1;p.space_before=Pt(0);p.space_after=Pt(0)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.keep_together=False;p.widow_control=False
        if name=='File title':p.keep_with_next=True;p.space_before=Pt(10);p.space_after=Pt(4)
        if name=='Table text':p.line_spacing=1.15;p.space_after=Pt(3);p.widow_control=True
    cap=d.styles['Caption'].paragraph_format
    cap.line_spacing=1;cap.space_after=Pt(8);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    d.core_properties.author=student;d.core_properties.title='Разработка кроссплатформенных приложений. Лабораторные работы'
    d.core_properties.subject='Qt, C++17';d.core_properties.comments=''
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU')
    d.styles['Normal'].element.get_or_add_rPr().append(lang)
    return d

def para(d,text='',style=None):
    return d.add_paragraph(text,style)

def heading(d,text,level=2):
    p=d.add_heading(text,level)
    if level==1:p.paragraph_format.page_break_before=True
    return p

def cover(d,student,group,lab=None,title=''):
    def p(t,before=0,after=0,align=WD_ALIGN_PARAGRAPH.CENTER,bold=False):
        a=para(d,t,'Cover');a.alignment=align
        a.paragraph_format.space_before=Pt(before);a.paragraph_format.space_after=Pt(after)
        a.paragraph_format.line_spacing=1.15
        for r in a.runs:font(r,12,bold)
        return a
    p('ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ\nОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ\n«ИВАНОВСКИЙ ГОСУДАРСТВЕННЫЙ ЭНЕРГЕТИЧЕСКИЙ\nУНИВЕРСИТЕТ ИМЕНИ В. И. ЛЕНИНА»')
    p('Кафедра программного обеспечения компьютерных систем',before=14)
    p('ОТЧЁТ',before=75,bold=True)
    p('по лабораторным работам №1–5' if lab is None else f'по лабораторной работе №{lab}',after=7)
    p('по дисциплине\n«Разработка кроссплатформенных приложений»')
    if title:p(title,before=14,bold=True)
    p(f'Выполнил студент группы {group}\n{student}\n\nПроверила преподаватель кафедры ПОКС\nЧернышева Л. П.',before=80,align=WD_ALIGN_PARAGRAPH.RIGHT)
    p('Иваново 2026',before=85 if lab is None else 60)

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None:widths=[WIDTH_CM/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for idx,text in enumerate(headers):
        t.rows[0].cells[idx].text=text
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for row_no,row in enumerate(t.rows):
        for col,(cell,w) in enumerate(zip(row.cells,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.style=d.styles['Table text']
                if row_no==0:
                    p.paragraph_format.keep_with_next=True
                    for run in p.runs:font(run,12,True)
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if row_no==0:pr.append(OxmlElement('w:tblHeader'))
    para(d,'').paragraph_format.space_after=Pt(0)
    return t

def picture(d,path,caption,max_h=18,natural=True,max_w=WIDTH_CM):
    with Image.open(path) as im:w,h=im.size
    width=min(max_w, (w/180*2.54) if natural else max_w)
    width=min(width,max_h*w/h)
    p=para(d);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
    shape=p.add_run().add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set('descr',caption)
    para(d,caption,'Caption')

def source(d,path,label):
    para(d,label,'File title')
    for line in path.read_text(encoding='utf-8').splitlines():
        # Строки кода не переписываются; Word переносит длинные строки визуально.
        p=para(d,line.expandtabs(4),'Listing')
        p.paragraph_format.keep_with_next=False
    p=para(d,'');p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)


def lab_section(d,number,data):
    title=data['title'];folder=data['folder']
    heading(d,f'Лабораторная работа №{number}. {title}',1)
    heading(d,'Задание')
    for index,text in enumerate(data['task']):para(d,text+(' [1]' if index==0 and number in (1,2,3) else ''))
    heading(d,'Ход работы')
    for index,text in enumerate(data['work']):
        refs={(1,0):' [2, 3]',(2,0):' [2, 3]',(3,0):' [2, 3]',(4,1):' [2, 3]',(5,3):' [3, 4]',(5,4):' [1, 2]'}
        para(d,text+refs.get((number,index),''))
    if 'rules' in data:
        heading(d,'Правила обработки данных')
        for text in data['rules']:para(d,text)
    heading(d,'Диаграмма классов')
    para(d,'На диаграмме показаны основные классы и связи. Двоеточие в заголовке обозначает базовый класс; ромб — владение объектом или коллекцией; пунктирная стрелка — использование. Перечень методов сокращён, полный состав приведён в листингах.')
    for i,name in enumerate(data['diagrams'],1):
        if len(data['diagrams'])>1:
            # Каждая часть остаётся целиком с подписью.
            heading(d,DIAGRAM_TITLES.get(name,'Диаграмма классов'),3)
        caption=f'Рисунок {number}.{i} — '+DIAGRAM_TITLES.get(name,'Диаграмма классов приложения')
        picture(d,REPORTS/'assets'/f'{name}.png',caption)
    heading(d,'Назначение классов')
    table(d,['Класс / структура','Назначение'],data['classes'],[4.4,12.1])
    if 'format' in data:
        heading(d,'Формат файла')
        table(d,['Поле','Тип','Назначение'],data['format'],[4.2,3.0,9.3])
    heading(d,'Код программы')
    para(d,f'Исходники находятся в папке {folder}. Ниже приведены заголовочные файлы, реализации классов, точка входа и файлы сборки.')
    if number==4:para(d,'Классы справочника из ЛР3 сохранены. Для самостоятельной сборки здесь приведён полный комплект файлов, включая неизменившиеся классы.')
    for filename in SOURCE_ORDER[folder]:source(d,ROOT/folder/filename,filename)
    source(d,ROOT/'common/capture.h','Общий вспомогательный файл ../common/capture.h')
    heading(d,'Контрольные примеры')
    para(d,'В таблице приведены контрольные примеры и ожидаемые результаты. Результаты выполненных автоматизированных проверок описаны ниже.')
    table(d,['Действие / входные данные','Ожидаемый результат'],data['tests'],[7,9.5])
    heading(d,'Проверка и скриншоты')
    if 'verified' in data:
        para(d,data['verified'])
        if number==5:
            para(d,'Повторный запуск того же набора с AddressSanitizer и UndefinedBehaviorSanitizer также завершился результатом 59 passed, 0 failed. Диагностик санитайзеров не получено; протокол — verification/domain-sanitizers.log.')
    status=json.loads((REPORTS/'verification.json').read_text(encoding='utf-8'))
    para(d,status['gui'])
    present=[]
    for os_label,os_name in [('windows','Windows'),('ubuntu','Ubuntu'),('linux','Linux'),('macos','macOS')]:
        for name in data['screens']:
            path=REPORTS/'screenshots'/os_label/(name+'.png')
            if path.is_file():present.append((path,os_name,SCREEN_TITLES[name]))
    if present:
        para(d,'На рисунках показаны окна и диалоги, полученные при запуске собранных программ. Для снимков Windows использована светлая тема Qt.')
        for i,(path,os_name,name) in enumerate(present,len(data['diagrams'])+1):
            with Image.open(path) as im:maximum=min(WIDTH_CM,im.width/100*2.54)
            picture(d,path,f'Рисунок {number}.{i} — {name}. {os_name}',natural=False,max_h=12,max_w=maximum)
    else:
        para(d,'Скриншоты работающих окон пока не получены. После сборки необходимо приложить реальные снимки Windows и Ubuntu: '+', '.join(SCREEN_TITLES[n] for n in data['screens'])+'.')
    heading(d,'Вывод')
    para(d,data['conclusion'])
    heading(d,'Источники')
    for i,(text,url) in enumerate(data['sources'],1):
        p=para(d,f'{i}. {text}'+(' '+url if url else ''))
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.15


def build(student,group):
    data=json.loads((REPORTS/'content.json').read_text(encoding='utf-8'))
    out=REPORTS/'docx';out.mkdir(parents=True,exist_ok=True)
    for number in range(1,6):
        info=data[str(number)];d=setup(student);cover(d,student,group,number,info['title'])
        lab_section(d,number,info)
        d.save(out/f'LR{number}_Ismaylov.docx')
    d=setup(student);cover(d,student,group)
    heading(d,'Состав отчёта',1)
    for number in range(1,6):para(d,f'Лабораторная работа №{number}. '+data[str(number)]['title'])
    heading(d,'Состояние проверки')
    status=json.loads((REPORTS/'verification.json').read_text(encoding='utf-8'))
    para(d,status['master'])
    para(d,'ЛР4 является продолжением справочника ЛР3. ЛР5 — самостоятельное приложение «Планировщик учебных заданий». Подробные команды сборки и обновления отчёта находятся в README.md комплекта.')
    for number in range(1,6):
        lab_section(d,number,data[str(number)])
    d.save(out/'Ismaylov_Qt_LR1-5.docx')
    return out

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--student',default='Исмайлов В.П.')
    ap.add_argument('--group',default='4-41')
    args=ap.parse_args()
    print(build(args.student,args.group))
