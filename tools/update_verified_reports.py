"""Добавить проверенные результаты и снимки в существующие отчёты без смены стилей."""
import json
import re
from pathlib import Path
from docx import Document
from PIL import Image
from build_reports import picture, SCREEN_TITLES

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / 'reports'
ENVIRONMENT = ('Приложения собраны и запущены в Windows 11 с Qt 6.9.2 и GCC 15.2.0 '
               'и в Ubuntu 24.04 с Qt 6.4.2 и GCC 13.3.0. В Ubuntu использован '
               'виртуальный экран X11 Xvfb в контейнере Docker. Проверка выполнена 6 октября 2026 г.')
MASTER = ('Все пять приложений собраны и запущены в Windows 11 и Ubuntu 24.04. '
          'В каждой системе прошли 59 проверок ядра StudyPlan и 15 сценариев Qt Test. '
          'Результат Qt Test — 17 passed, 0 failed, включая начальную и заключительную '
          'служебные проверки. В отчёт включены 24 снимка работающих окон и диалогов. '
          'Ubuntu проверена в контейнере с виртуальным экраном X11 Xvfb.')
RESULTS = {
 1: 'Проверены арифметические операции, дробные и отрицательные числа, ввод с клавиатуры, деление на ноль, восстановление после ошибки и повторный расчёт с результатом в научной форме. Ожидаемые значения совпали с полученными.',
 2: 'Проверены диапазоны часов и минут и исходное выключенное состояние. При установке текущих часа и минуты появилось сообщение «Пора вставать!». Повторный вызов обновления времени в ту же минуту не создал второго сообщения.',
 3: 'Проверены модель данных, изменение названия, отказ при пустом названии и недопустимых числах, очистка коллекции и удаление при отсутствии записей. QAbstractItemModelTester не обнаружил нарушений контракта модели.',
 4: 'Через диалог добавлена марка «Исследование космоса». Отмена закрытия сохранила открытое окно и изменения. Коллекция записана в файл и открыта в новом окне; состав и поля записей совпали. Повреждённый файл не заменил прежнюю коллекцию, а запись по пути каталога завершилась ошибкой.',
 5: 'Через диалог создано задание «Подготовка к защите» и записано занятие длительностью 45 минут. Задание перешло в состояние «В работе». После отметки о выполнении сводка показала 100 %. Повторная загрузка сохранила задание, занятие, статус и фактическое время. Проверены резервное копирование, восстановление связей, фильтрация, JSON, отказ записи и экранирование CSV. Все 59 проверок ядра и 15 сценариев Qt Test прошли в обеих системах.'
}
SCREENS = {
 1: ['lab1'], 2: ['lab2', 'lab2_signal'], 3: ['lab3'],
 4: ['lab4', 'lab4_dialog', 'lab4_saved'],
 5: ['lab5_tasks', 'lab5_dialog', 'lab5_sessions', 'lab5_summary', 'lab5_progress']
}
SCREEN_TITLES.update({
 'lab2_signal': 'Сообщение о срабатывании будильника',
 'lab4_dialog': 'Добавление марки', 'lab4_saved': 'Коллекция после сохранения и повторного открытия',
 'lab5_dialog': 'Добавление учебного задания', 'lab5_progress': 'Сводка после занятия и выполнения задания'
})

def append_before(document, anchor, text):
    p = document.add_paragraph(text)
    anchor.addprevious(p._p)

def update(path, data):
    document = Document(path)
    for table in document.tables:
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    section = 0
    paragraphs = list(document.paragraphs)
    for index, p in enumerate(paragraphs):
        match = re.match(r'Лабораторная работа №(\d)', p.text)
        if match and p.style.name == 'Heading 1': section = int(match[1])
        if p.text.startswith('В отчёте приведены исходники и описание пяти лабораторных работ.'):
            p.text = MASTER
        if p.text.startswith('В таблице приведены контрольные примеры и ожидаемые результаты.'):
            p.text = 'В таблице приведены контрольные примеры и ожидаемые результаты. Результаты выполненных автоматизированных проверок описаны ниже.'
        if 'Графическая часть и сериализация Qt подготовлены к проверке' in p.text:
            p.text = p.text.replace('Графическая часть и сериализация Qt подготовлены к проверке в целевых операционных системах.', 'Графическая часть и сериализация Qt проверены в Windows 11 и Ubuntu 24.04.')
        if p.text != 'Проверка и скриншоты': continue
        end = index + 1
        while end < len(paragraphs) and paragraphs[end].text != 'Вывод': end += 1
        if end == len(paragraphs): raise RuntimeError('Не найден раздел Вывод')
        for old in paragraphs[index+1:end]:
            old._p.getparent().remove(old._p)
        anchor = paragraphs[end]._p
        append_before(document, anchor, ENVIRONMENT)
        append_before(document, anchor, RESULTS[section])
        if section == 5:
            append_before(document, anchor, 'Дополнительная проверка ядра в Debian GNU/Linux 13 с AddressSanitizer и UndefinedBehaviorSanitizer также завершилась без ошибок. Протоколы всех запусков находятся в папке verification.')
        append_before(document, anchor, 'На рисунках показаны окна и диалоги, полученные при запуске собранных программ. Для снимков Windows использована светлая тема Qt.')
        number = len(data[str(section)]['diagrams']) + 1
        for os_label, os_name in [('windows', 'Windows 11'), ('ubuntu', 'Ubuntu 24.04')]:
            for screen in SCREENS[section]:
                image = REPORTS / 'screenshots' / os_label / (screen+'.png')
                if not image.is_file(): raise FileNotFoundError(image)
                before = len(document.paragraphs)
                with Image.open(image) as im: maximum=min(16.5,im.width/100*2.54)
                picture(document, image, f'Рисунок {section}.{number} — {SCREEN_TITLES[screen]}. {os_name}', max_h=12, natural=False,max_w=maximum)
                for added in document.paragraphs[before:]: anchor.addprevious(added._p)
                number += 1
    document.save(path)

if __name__ == '__main__':
    data = json.loads((REPORTS/'content.json').read_text(encoding='utf-8'))
    data['5']['conclusion'] = data['5']['conclusion'].replace('Графическая часть и сериализация Qt подготовлены к проверке в целевых операционных системах.', 'Графическая часть и сериализация Qt проверены в Windows 11 и Ubuntu 24.04.')
    for path in sorted((REPORTS/'docx').glob('*.docx')):
        update(path, data)
        print(path.name)
    for number in range(1,6):
        data[str(number)]['screens'] = SCREENS[number]
        data[str(number)]['verified'] = RESULTS[number]
    (REPORTS/'content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (REPORTS/'verification.json').write_text(json.dumps({'gui':ENVIRONMENT,'master':MASTER},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
